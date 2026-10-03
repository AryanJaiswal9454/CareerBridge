
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


TEMPLATES = [
    {
        "key": "aryan_ats",
        "name": "Aryan ATS",
        "description": "Based on the uploaded one-column ATS-friendly resume format.",
        "recommended": True,
    },
    {
        "key": "classic_ats",
        "name": "Classic ATS",
        "description": "Traditional professional single-column format.",
        "recommended": False,
    },
    {
        "key": "modern_professional",
        "name": "Modern Professional",
        "description": "Clean modern layout while remaining ATS readable.",
        "recommended": False,
    },
    {
        "key": "compact_ats",
        "name": "Compact ATS",
        "description": "Space-efficient format for one-page resumes.",
        "recommended": False,
    },
]


def get_templates():
    return TEMPLATES


def _set_font(run, name="Arial", size=9.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold


def _add_border(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()

    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")

    pBdr.append(bottom)
    pPr.append(pBdr)


def _heading(document, title, size=10.5):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(title.upper())
    _set_font(run, size=size, bold=True)

    _add_border(paragraph)
    return paragraph


def _bullet(document, text, size=9.2):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)

    run = paragraph.add_run(str(text))
    _set_font(run, size=size)

    return paragraph


def _setup_document(
    top=0.55,
    bottom=0.55,
    left=0.65,
    right=0.65,
):
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(2)

    return document


def _header(document, data, name_size=18):
    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)

    run = name.add_run(data.get("full_name") or "RESUME")
    _set_font(run, size=name_size, bold=True)

    contact = data.get("contact_line") or ""

    if contact:
        contact_p = document.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_p.paragraph_format.space_after = Pt(6)

        run = contact_p.add_run(contact)
        _set_font(run, size=9)


def _summary(document, data):
    if data.get("professional_summary"):
        _heading(document, "Professional Summary")
        paragraph = document.add_paragraph(data["professional_summary"])
        paragraph.paragraph_format.space_after = Pt(2)


def _skills(document, data):
    skills = data.get("key_skills") or []

    if not skills:
        return

    _heading(document, "Technical Skills")

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(" | ".join(str(x) for x in skills))
    _set_font(run)


def _experience(document, data):
    jobs = data.get("work_experience") or []

    if not jobs:
        return

    _heading(document, "Professional Experience")

    for job in jobs:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)

        title = job.get("title", "")
        organization = job.get("organization", "")
        duration = job.get("duration", "")

        run = paragraph.add_run(
            f"{title} — {organization}".strip(" —")
        )
        _set_font(run, size=10, bold=True)

        if duration:
            run = paragraph.add_run(f" | {duration}")
            _set_font(run, size=9.2)

        for bullet in job.get("bullets", []):
            _bullet(document, bullet)


def _projects(document, data):
    projects = data.get("projects") or []

    if not projects:
        return

    _heading(document, "Projects")

    for project in projects:
        if isinstance(project, dict):
            name = project.get("name") or project.get("title") or "Project"
            description = project.get("description") or ""
            tech_stack = project.get("tech_stack") or ""
            bullets = project.get("bullets") or []

            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(1)

            run = paragraph.add_run(name)
            _set_font(run, size=10, bold=True)

            if tech_stack:
                run = paragraph.add_run(f" | {tech_stack}")
                _set_font(run, size=9.2)

            if description:
                _bullet(document, description)

            for bullet in bullets:
                _bullet(document, bullet)

        else:
            _bullet(document, project)


def _education(document, data):
    education = data.get("education") or []

    if not education:
        return

    _heading(document, "Education")

    for item in education:
        _bullet(document, item)


def _certifications(document, data):
    certifications = data.get("certifications") or []

    if not certifications:
        return

    _heading(document, "Certifications")

    for item in certifications:
        _bullet(document, item)


def _build_common(
    data,
    margins,
    name_size,
):
    document = _setup_document(*margins)

    _header(document, data, name_size=name_size)
    _summary(document, data)
    _skills(document, data)
    _experience(document, data)
    _projects(document, data)
    _education(document, data)
    _certifications(document, data)

    return document


def build_aryan_ats(data):
    """
    Template 1:
    Based on the uploaded resume format.
    """
    return _build_common(
        data,
        margins=(0.55, 0.55, 0.65, 0.65),
        name_size=18,
    )


def build_classic_ats(data):
    document = _build_common(
        data,
        margins=(0.65, 0.65, 0.75, 0.75),
        name_size=17,
    )
    return document


def build_modern_professional(data):
    document = _build_common(
        data,
        margins=(0.60, 0.60, 0.70, 0.70),
        name_size=19,
    )

    # Slightly larger section spacing.
    for paragraph in document.paragraphs:
        if paragraph.text.isupper():
            paragraph.paragraph_format.space_before = Pt(9)

    return document


def build_compact_ats(data):
    document = _build_common(
        data,
        margins=(0.45, 0.45, 0.55, 0.55),
        name_size=16,
    )

    normal = document.styles["Normal"]
    normal.font.size = Pt(9)

    return document


BUILDERS = {
    "aryan_ats": build_aryan_ats,
    "classic_ats": build_classic_ats,
    "modern_professional": build_modern_professional,
    "compact_ats": build_compact_ats,
}


def build_resume_docx(data, template_key="aryan_ats"):
    builder = BUILDERS.get(template_key)

    if not builder:
        template_key = "aryan_ats"
        builder = BUILDERS[template_key]

    document = builder(data)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer
