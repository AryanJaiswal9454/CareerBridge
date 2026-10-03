from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


TEMPLATE_META = [
    {
        "key": "aryan_ats",
        "name": "Aryan ATS",
        "description": "Clean one-column ATS-friendly format based on the master resume style.",
        "recommended": True,
    },
    {
        "key": "classic_ats",
        "name": "Classic ATS",
        "description": "Traditional professional single-column resume layout.",
        "recommended": False,
    },
    {
        "key": "modern_professional",
        "name": "Modern Professional",
        "description": "Clean modern typography while remaining ATS readable.",
        "recommended": False,
    },
    {
        "key": "compact_ats",
        "name": "Compact ATS",
        "description": "Space-efficient layout for concise one-page resumes.",
        "recommended": False,
    },
]


def _text(value):
    return str(value or "").strip()


def _list(value):
    if not value:
        return []
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict) or _text(item)]
    return [_text(value)]


def _add_bottom_border(paragraph, color="64748B"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _set_run(run, size=9.5, bold=False, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _heading(document, title, style="underline"):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(_text(title).upper())
    _set_run(run, size=10, bold=True)
    if style == "underline":
        _add_bottom_border(paragraph)
    return paragraph


def _bullet(document, value, compact=False):
    value = _text(value)
    if not value:
        return
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(1 if compact else 2)
    run = paragraph.add_run(value)
    _set_run(run, size=8.8 if compact else 9.2)


def _setup_document(template_key):
    document = Document()
    section = document.sections[0]

    margins = {
        "aryan_ats": (0.55, 0.55, 0.65, 0.65),
        "classic_ats": (0.65, 0.65, 0.75, 0.75),
        "modern_professional": (0.6, 0.6, 0.72, 0.72),
        "compact_ats": (0.45, 0.45, 0.55, 0.55),
    }.get(template_key, (0.55, 0.55, 0.65, 0.65))

    section.top_margin = Inches(margins[0])
    section.bottom_margin = Inches(margins[1])
    section.left_margin = Inches(margins[2])
    section.right_margin = Inches(margins[3])

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.2 if template_key == "compact_ats" else 9.5)
    normal.paragraph_format.space_after = Pt(2 if template_key == "compact_ats" else 3)

    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.size = Pt(10.5)

    return document


def _contact_line(data):
    values = []
    for key in ("location", "phone", "email", "linkedin", "github", "portfolio"):
        value = _text(data.get(key))
        if value:
            values.append(value)
    return " | ".join(values)


def _add_header(document, data, template_key):
    name = document.add_paragraph()
    name.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if template_key in ("classic_ats", "modern_professional")
        else WD_ALIGN_PARAGRAPH.CENTER
    )
    name.paragraph_format.space_after = Pt(1)
    run = name.add_run(_text(data.get("full_name")) or "YOUR NAME")
    _set_run(
        run,
        size={
            "aryan_ats": 18,
            "classic_ats": 17,
            "modern_professional": 19,
            "compact_ats": 16,
        }.get(template_key, 18),
        bold=True,
    )

    headline = _text(data.get("headline"))
    if headline:
        p = document.add_paragraph()
        p.alignment = name.alignment
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(headline)
        _set_run(run, size=9.5, bold=True, color="0891B2" if template_key == "modern_professional" else None)

    contact = _contact_line(data)
    if contact:
        p = document.add_paragraph()
        p.alignment = name.alignment
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(contact)
        _set_run(run, size=8.5)


def _add_summary(document, data):
    summary = _text(data.get("professional_summary"))
    if summary:
        _heading(document, "Professional Summary")
        document.add_paragraph(summary)


def _add_skills(document, data):
    skills = data.get("key_skills") or []
    if skills:
        _heading(document, "Technical Skills")
        document.add_paragraph(" | ".join(_text(x) for x in skills if _text(x)))


def _add_experience(document, data, compact=False):
    jobs = [x for x in (data.get("work_experience") or []) if isinstance(x, dict)]
    jobs = [x for x in jobs if any(_text(x.get(k)) for k in ("title", "organization", "bullets"))]
    if not jobs:
        return

    _heading(document, "Professional Experience")
    for job in jobs:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        title = _text(job.get("title"))
        organization = _text(job.get("organization"))
        duration = _text(job.get("duration"))
        line = " — ".join(x for x in (title, organization) if x)
        if duration:
            line = f"{line} | {duration}" if line else duration
        run = p.add_run(line)
        _set_run(run, size=9.3 if compact else 9.8, bold=True)
        for bullet in job.get("bullets") or []:
            _bullet(document, bullet, compact=compact)


def _add_projects(document, data, compact=False):
    projects = [x for x in (data.get("projects") or []) if isinstance(x, dict)]
    projects = [x for x in projects if any(_text(x.get(k)) for k in ("name", "description", "bullets", "tech_stack"))]
    if not projects:
        return

    _heading(document, "Projects")
    for project in projects:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        name = _text(project.get("name"))
        run = p.add_run(name)
        _set_run(run, size=9.3 if compact else 9.8, bold=True)
        if _text(project.get("description")):
            p = document.add_paragraph(_text(project.get("description")))
            p.paragraph_format.space_after = Pt(1)
        for bullet in project.get("bullets") or []:
            _bullet(document, bullet, compact=compact)
        tech = _text(project.get("tech_stack"))
        if tech:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(f"Tech stack: {tech}")
            _set_run(run, size=8.5 if compact else 8.8)


def _add_education(document, data, compact=False):
    education = [x for x in (data.get("education") or []) if isinstance(x, dict)]
    if not education:
        return

    _heading(document, "Education")
    for item in education:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        degree = _text(item.get("degree"))
        institution = _text(item.get("institution"))
        duration = _text(item.get("duration"))
        grade = _text(item.get("grade"))
        line = " — ".join(x for x in (degree, institution) if x)
        if duration:
            line = f"{line} | {duration}" if line else duration
        run = p.add_run(line)
        _set_run(run, size=9 if compact else 9.3, bold=True)
        if grade:
            p2 = document.add_paragraph(grade)
            p2.paragraph_format.space_after = Pt(1)


def _add_certifications(document, data, compact=False):
    certifications = [x for x in (data.get("certifications") or []) if isinstance(x, dict)]
    if not certifications:
        return

    _heading(document, "Certifications")
    for cert in certifications:
        name = _text(cert.get("name"))
        issuer = _text(cert.get("issuer"))
        date = _text(cert.get("date"))
        line = " — ".join(x for x in (name, issuer) if x)
        if date:
            line = f"{line} | {date}" if line else date
        _bullet(document, line, compact=compact)


def build_resume_docx(data, template_key="aryan_ats"):
    template_key = template_key if template_key in {x["key"] for x in TEMPLATE_META} else "aryan_ats"
    compact = template_key == "compact_ats"
    document = _setup_document(template_key)

    _add_header(document, data, template_key)
    _add_summary(document, data)
    _add_skills(document, data)
    _add_experience(document, data, compact)
    _add_projects(document, data, compact)
    _add_education(document, data, compact)
    _add_certifications(document, data, compact)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def resume_to_text(data):
    parts = []
    for key in ("full_name", "headline", "location", "phone", "email", "linkedin", "github", "portfolio", "professional_summary"):
        value = _text(data.get(key))
        if value:
            parts.append(value)

    parts.extend(_text(x) for x in (data.get("key_skills") or []) if _text(x))

    for job in data.get("work_experience") or []:
        if not isinstance(job, dict):
            continue
        parts.extend(_text(job.get(k)) for k in ("title", "organization", "duration") if _text(job.get(k)))
        parts.extend(_text(x) for x in (job.get("bullets") or []) if _text(x))

    for project in data.get("projects") or []:
        if not isinstance(project, dict):
            continue
        parts.extend(_text(project.get(k)) for k in ("name", "description", "tech_stack") if _text(project.get(k)))
        parts.extend(_text(x) for x in (project.get("bullets") or []) if _text(x))

    for item in data.get("education") or []:
        if not isinstance(item, dict):
            continue
        parts.extend(_text(item.get(k)) for k in ("degree", "institution", "duration", "grade") if _text(item.get(k)))

    for cert in data.get("certifications") or []:
        if not isinstance(cert, dict):
            continue
        parts.extend(_text(cert.get(k)) for k in ("name", "issuer", "date") if _text(cert.get(k)))

    return "\n".join(x for x in parts if x)
